from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from openpyxl import Workbook

from ToolApp.models import Users
from ToolApp.salary_import import SalarySourceError, parse_salary_files


class ImportEmployeeSalariesCommandTests(TestCase):
    def setUp(self):
        self.ion = Users.objects.create(
            UserName="Ion Popescu U1234567",
            UserSerie="1234567890123_U1234567_CIM 01.01.2026",
            Company="DMX",
            person_type=Users.PersonType.EMPLOYEE,
        )
        self.xmeg = Users.objects.create(
            UserName="David Gheorghe Dumitru // notă internă",
            UserSerie="XMEG-1",
            Company="XMEG CONSTRUCT",
            person_type=Users.PersonType.EMPLOYEE,
        )
        Users.objects.create(
            UserName="Davinder Singh R5213071",
            UserSerie="7850722320031_R5213071",
            Company="VB-ROM",
            person_type=Users.PersonType.EMPLOYEE,
        )

    def make_workbook(self, directory):
        path = Path(directory) / "salarii-test.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "DMX"
        sheet.append([None, None, None, "SALARII IULIE 2026"])
        sheet.append([None, None, "NR. CRT.", "NUME/PRENUME ANGAJAT"])
        sheet.append([
            None, None, "DMX CONSTRUCTION", None, "Avans Salar", "Poprire",
            "Lichidare", "Tichete", "net+tichete stat plata",
        ])
        sheet.append([None, None, 1, "POPESCU ION_1234567890123_U1234567_CIM", 1000, 200, 3000, 400, 4600])
        sheet.append([None, None, 2, "NUME NEGĂSIT_CIM", 500, None, 900, 100, 1500])
        # Același angajat pe încă un rând trebuie agregat, nu suprascris.
        sheet.append([None, None, 3, "ION POPESCU_1234567890123", 100, None, 200, 0, 300])
        sheet.append([None, None, 4, "DALVIR SINGH", 100, None, 200, 0, 300])

        xmeg = workbook.create_sheet("XMEG")
        xmeg.append([])
        xmeg.append([None, None, None, "XMEG", None, "AVANS STAT PLATA", "DIF", None, "lichidare", "tichete", "total stat plata"])
        xmeg.append([None, None, None, "DAVID GH DUMITRU", 1800, 534, 1266, None, 2741, 760, 4035])
        workbook.save(path)
        return path

    def test_dry_run_lists_unmatched_and_does_not_change_database(self):
        with TemporaryDirectory() as directory:
            path = self.make_workbook(directory)
            output = StringIO()
            call_command("import_employee_salaries", str(path), stdout=output)

        self.ion.refresh_from_db()
        self.assertIsNone(self.ion.total_salary_ron)
        self.assertIn("Nume negăsite: 2", output.getvalue())
        self.assertIn("NUME NEGĂSIT", output.getvalue())
        self.assertIn("DALVIR SINGH", output.getvalue())
        self.assertIn("DRY-RUN", output.getvalue())

    def test_apply_imports_fields_adds_garnishment_and_uses_labeled_xmeg_advance(self):
        with TemporaryDirectory() as directory:
            path = self.make_workbook(directory)
            output = StringIO()
            call_command("import_employee_salaries", str(path), "--apply", stdout=output)

        self.ion.refresh_from_db()
        self.assertEqual(str(self.ion.total_salary_ron), "4900.00")
        self.assertEqual(str(self.ion.salary_advance_ron), "1100.00")
        self.assertEqual(str(self.ion.salary_remainder_ron), "3400.00")
        self.assertEqual(str(self.ion.meal_vouchers_ron), "400.00")

        self.xmeg.refresh_from_db()
        self.assertEqual(str(self.xmeg.total_salary_ron), "4035.00")
        self.assertEqual(str(self.xmeg.salary_advance_ron), "534.00")
        self.assertEqual(str(self.xmeg.salary_remainder_ron), "2741.00")
        self.assertEqual(str(self.xmeg.meal_vouchers_ron), "760.00")
        self.assertIn("coloana DIF (1266.00 lei) a fost ignorată", output.getvalue())
        self.assertIn("Import finalizat: 2 angajați actualizați", output.getvalue())


class ConsolidatedSalaryImportTests(TestCase):
    def setUp(self):
        self.employee = Users.objects.create(
            UserName="Ion Popescu", UserSerie="1234567890123_U1234567_CIM",
            Company="DMX", total_salary_ron="9000", salary_remainder_ron="8000",
            salary_advance_ron="1000", meal_vouchers_ron="300",
            hourly_rate="40", bonus_lunar_sef_echipa="150",
        )
        self.outside = Users.objects.create(
            UserName="Alt Angajat", UserSerie="outside", total_salary_ron="7000",
        )

    def workbook(self, directory, rows):
        path = Path(directory) / "consolidat.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Salarii"
        sheet.append(["Salarii iulie 2026"])
        sheet.append(["Nume angajat", "Avans (lei)", "Bonuri (lei)", "Lichidare (lei)", "Total (lei)", "Firmă", "Pașaport", "Situație"])
        for row in rows:
            sheet.append(row)
        sheet.append(["TOTAL", 99999, 99999, 99999, 99999])
        other = book.create_sheet("Lipsa din copie")
        other.append(["Nume angajat", "Firmă", "Avans (lei)", "Bonuri (lei)"])
        other.append(["Ion Popescu", "DMX", 1100, 400])
        book.save(path)
        return path

    def run_import(self, path, apply=True):
        output = StringIO()
        call_command("import_employee_salaries", str(path), apply=apply, require_matches=True, stdout=output)
        return output.getvalue()

    def test_full_import_is_idempotent_and_only_updates_salary_fields(self):
        with TemporaryDirectory() as directory:
            path = self.workbook(directory, [["Different display name", 1100, 400, 3500, 5000, "DMX", "U1234567"]])
            self.run_import(path)
            self.run_import(path)
        self.employee.refresh_from_db()
        self.outside.refresh_from_db()
        self.assertEqual(str(self.employee.total_salary_ron), "5000.00")
        self.assertEqual(str(self.employee.salary_advance_ron), "1100.00")
        self.assertEqual(str(self.employee.salary_remainder_ron), "3500.00")
        self.assertEqual(str(self.employee.meal_vouchers_ron), "400.00")
        self.assertEqual(str(self.employee.hourly_rate), "40.00")
        self.assertEqual(str(self.employee.bonus_lunar_sef_echipa), "150.00")
        self.assertEqual(str(self.outside.total_salary_ron), "7000.00")
        self.assertEqual(Users.objects.count(), 2)

    def test_missing_totals_preserve_existing_values_and_blank_vouchers_mean_zero(self):
        with TemporaryDirectory() as directory:
            path = self.workbook(directory, [["POPESCU ION", 1200, None, None, None, "DMX", None]])
            output = self.run_import(path)
        self.employee.refresh_from_db()
        self.assertEqual(str(self.employee.total_salary_ron), "9000.00")
        self.assertEqual(str(self.employee.salary_remainder_ron), "8000.00")
        self.assertEqual(str(self.employee.salary_advance_ron), "1200.00")
        self.assertEqual(str(self.employee.meal_vouchers_ron), "0.00")
        self.assertIn("Rânduri salariale citite: 1", output)

    def test_dry_run_preserves_database(self):
        with TemporaryDirectory() as directory:
            path = self.workbook(directory, [["Ion Popescu", 1100, 400, 3500, 5000, "DMX", "U1234567"]])
            self.run_import(path, apply=False)
        self.employee.refresh_from_db()
        self.assertEqual(str(self.employee.total_salary_ron), "9000.00")

    def test_ambiguous_names_and_wrong_passports_are_not_updated(self):
        Users.objects.create(UserName="Ion Popescu", UserSerie="second", Company="DMX")
        with TemporaryDirectory() as directory:
            path = self.workbook(directory, [
                ["Ion Popescu", 1100, 400, 3500, 5000, "DMX", None],
                ["Ion Popescu", 1100, 400, 3500, 5000, "DMX", "U7654321"],
            ])
            with self.assertRaisesMessage(CommandError, "Niciun angajat asociat"):
                self.run_import(path)
        self.employee.refresh_from_db()
        self.assertEqual(str(self.employee.total_salary_ron), "9000.00")

    def test_invalid_balance_aborts_before_any_write(self):
        with TemporaryDirectory() as directory:
            path = self.workbook(directory, [
                ["Ion Popescu", 1100, 400, 3500, 5000, "DMX", "U1234567"],
                ["Other", 1100, 400, 3500, 6000, "DMX", None],
            ])
            with self.assertRaisesMessage(CommandError, "Total nu este egal"):
                self.run_import(path)
        self.employee.refresh_from_db()
        self.assertEqual(str(self.employee.total_salary_ron), "9000.00")

    def test_formula_without_cached_value_is_rejected(self):
        with TemporaryDirectory() as directory:
            path = self.workbook(directory, [["Ion Popescu", 1100, 400, "=E3-B3-C3", 5000, "DMX", "U1234567"]])
            with self.assertRaisesMessage(SalarySourceError, "formulă fără rezultat salvat"):
                parse_salary_files([path])

    def test_mixed_complete_and_incomplete_rows_do_not_write_partial_total(self):
        with TemporaryDirectory() as directory:
            path = self.workbook(directory, [
                ["Ion Popescu", 1100, 400, 3500, 5000, "DMX", "U1234567"],
                ["Ion Popescu", 100, None, None, None, "DMX", "U1234567"],
            ])
            self.run_import(path)
        self.employee.refresh_from_db()
        self.assertEqual(str(self.employee.total_salary_ron), "9000.00")
        self.assertEqual(str(self.employee.salary_remainder_ron), "8000.00")
        self.assertEqual(str(self.employee.salary_advance_ron), "1200.00")

    def test_attached_workbook_has_170_rows_without_duplicate_report_or_totals(self):
        path = Path(__file__).parent / "data" / "Salarii_iulie_lichidare.xlsx"
        rows = parse_salary_files([path])
        self.assertEqual(len(rows), 170)
        self.assertEqual(sum(row.total is None for row in rows), 62)
        self.assertEqual(sum(row.remainder is None for row in rows), 62)
        self.assertTrue(all(row.sheet == "Salarii" for row in rows))
