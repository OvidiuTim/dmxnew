import io
from datetime import date, datetime

from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ToolApp.models import Histories, Tools
from ToolApp.security import request_has_admin


HEADER_FILL = PatternFill("solid", fgColor="123B2B")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ACCENT_FILL = PatternFill("solid", fgColor="DDF5E9")
DATE_FORMAT = "dd.mm.yyyy"


def _admin_required(request):
    return getattr(request, "dmx_role", None) == "admin" or request_has_admin(request)


def _parse_date(raw_value, label):
    raw_value = (raw_value or "").strip()
    if not raw_value:
        return None, None
    try:
        return date.fromisoformat(raw_value), None
    except ValueError:
        return None, JsonResponse(
            {"error": f"{label} trebuie să fie în formatul AAAA-LL-ZZ."},
            status=400,
        )


def _local_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        if timezone.is_aware(value):
            value = timezone.localtime(value)
        return value.date()
    return value


def _in_period(value, start_date, end_date):
    # Alocările inițiale fără dată rămân vizibile doar în raportul complet.
    if value is None:
        return start_date is None
    if start_date and value < start_date:
        return False
    return not end_date or value <= end_date


def _quantity(value):
    try:
        parsed = float(value or 1)
    except (TypeError, ValueError):
        return 1
    return int(parsed) if parsed.is_integer() else parsed


def tape_measure_report_rows(start_date=None, end_date=None):
    """Returnează toate predările de rulete și alocările inițiale încă active."""
    end_date = end_date or timezone.localdate()
    rows = []
    tools_with_recorded_out = set()

    histories = (
        Histories.objects
        .select_related("user_fk", "tool_fk", "issued_by")
        .filter(direction=Histories.Movement.OUT)
        .order_by("timestamp", "HistoryId")
    )
    for movement in histories:
        tool_name = movement.tool_fk.ToolName if movement.tool_fk_id else movement.Tool
        if "rulet" not in (tool_name or "").casefold():
            continue
        if movement.tool_fk_id:
            tools_with_recorded_out.add(movement.tool_fk_id)

        movement_date = movement.DateOfGiving or _local_date(movement.timestamp)
        if not _in_period(movement_date, start_date, end_date):
            continue

        employee = movement.user_fk
        recipient_name = employee.UserName if employee else (movement.User or "Nespecificat")
        rows.append({
            "event_id": f"history-{movement.HistoryId}",
            "recipient_id": employee.UserId if employee else None,
            "recipient_name": recipient_name,
            "recipient_series": employee.UserSerie if employee else "",
            "person_type": employee.get_person_type_display() if employee else "Nespecificat",
            "tool_name": tool_name or "Ruletă",
            "tool_series": movement.tool_fk.ToolSerie if movement.tool_fk_id else (movement.ToolSerie or ""),
            "quantity": _quantity(movement.quantity or movement.Pieces),
            "received_date": movement_date.isoformat() if movement_date else None,
            "issued_by": movement.issued_by.UserName if movement.issued_by_id else "",
            "source": "Istoric mișcări",
        })

    # Inventarele vechi au fost importate direct ca stare curentă și nu au
    # întotdeauna o mișcare OUT. Le includem o singură dată în raport.
    initial_allocations = (
        Tools.objects
        .select_related("AssignedTo")
        .filter(
            ToolName__icontains="rulet",
            AssignedTo__isnull=False,
            Status=Tools.ToolStatus.IN_LUCRU,
            IsReturned=False,
        )
        .order_by("DateOfGiving", "ToolId")
    )
    for tool in initial_allocations:
        if tool.ToolId in tools_with_recorded_out:
            continue
        received_date = tool.DateOfGiving
        if not _in_period(received_date, start_date, end_date):
            continue
        employee = tool.AssignedTo
        rows.append({
            "event_id": f"initial-{tool.ToolId}",
            "recipient_id": employee.UserId,
            "recipient_name": employee.UserName,
            "recipient_series": employee.UserSerie,
            "person_type": employee.get_person_type_display(),
            "tool_name": tool.ToolName,
            "tool_series": tool.ToolSerie or "",
            "quantity": _quantity(tool.Pieces),
            "received_date": received_date.isoformat() if received_date else None,
            "issued_by": "",
            "source": "Alocare inițială importată",
        })

    rows.sort(
        key=lambda item: (item["received_date"] or "0000-00-00", item["event_id"]),
        reverse=True,
    )
    return rows


def _summary(rows):
    recipients = {
        (f"id:{row['recipient_id']}" if row["recipient_id"] else f"name:{row['recipient_name'].casefold()}")
        for row in rows
    }
    dated_rows = [date.fromisoformat(row["received_date"]) for row in rows if row["received_date"]]
    return {
        "events": len(rows),
        "recipients": len(recipients),
        "quantity": sum(float(row["quantity"]) for row in rows),
        "first_date": min(dated_rows).isoformat() if dated_rows else None,
        "last_date": max(dated_rows).isoformat() if dated_rows else None,
    }


def _filter_rows(rows, raw_query):
    query = (raw_query or "").strip().casefold()
    if not query:
        return rows
    searchable_fields = (
        "recipient_name", "recipient_series", "tool_name", "tool_series", "issued_by",
    )
    return [
        row for row in rows
        if any(query in str(row[field] or "").casefold() for field in searchable_fields)
    ]


def _report_period(request):
    start_date, error = _parse_date(request.GET.get("start_date"), "Data de început")
    if error:
        return None, None, error
    end_date, error = _parse_date(request.GET.get("end_date"), "Data de sfârșit")
    if error:
        return None, None, error
    end_date = end_date or timezone.localdate()
    if start_date and start_date > end_date:
        return None, None, JsonResponse(
            {"error": "Data de început nu poate fi după data de sfârșit."},
            status=400,
        )
    return start_date, end_date, None


def tape_measure_report(request):
    if request.method != "GET":
        return JsonResponse({"error": "Metodă nepermisă."}, status=405)
    if not _admin_required(request):
        return JsonResponse({"error": "Nu ai dreptul să vezi raportul de rulete."}, status=403)
    start_date, end_date, error = _report_period(request)
    if error:
        return error
    rows = _filter_rows(
        tape_measure_report_rows(start_date, end_date),
        request.GET.get("q"),
    )
    return JsonResponse({
        "start_date": start_date.isoformat() if start_date else None,
        "end_date": end_date.isoformat(),
        "rows": rows,
        "summary": _summary(rows),
    })


def _safe_excel_text(value):
    value = "" if value is None else str(value)
    return f"'{value}" if value.lstrip().startswith(("=", "+", "-", "@")) else value


def _build_workbook(rows, start_date, end_date):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rulete predate"
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("A1:I1")
    sheet["A1"] = "Cine a primit rulete"
    sheet["A1"].font = Font(size=18, bold=True, color="123B2B")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    period_label = (
        f"Perioada: {start_date:%d.%m.%Y} – {end_date:%d.%m.%Y}"
        if start_date else f"Perioada: de la prima înregistrare – {end_date:%d.%m.%Y}"
    )
    sheet.merge_cells("A2:I2")
    sheet["A2"] = period_label
    sheet["A2"].font = Font(color="667085", italic=True)

    summary = _summary(rows)
    sheet["A4"] = "Predări"
    sheet["B4"] = summary["events"]
    sheet["D4"] = "Persoane distincte"
    sheet["E4"] = summary["recipients"]
    sheet["G4"] = "Bucăți"
    sheet["H4"] = summary["quantity"]
    for cell in ("A4", "D4", "G4"):
        sheet[cell].font = Font(bold=True, color="123B2B")
        sheet[cell].fill = ACCENT_FILL

    headers = [
        "Persoană", "Serie angajat", "Tip persoană", "Ruletă", "Serie unealtă",
        "Cantitate", "Data primirii", "Predat de", "Sursă înregistrare",
    ]
    header_row = 6
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for row_index, item in enumerate(rows, start=header_row + 1):
        received_date = date.fromisoformat(item["received_date"]) if item["received_date"] else None
        values = [
            _safe_excel_text(item["recipient_name"]),
            _safe_excel_text(item["recipient_series"]),
            item["person_type"],
            _safe_excel_text(item["tool_name"]),
            _safe_excel_text(item["tool_series"]),
            item["quantity"],
            received_date,
            _safe_excel_text(item["issued_by"]),
            item["source"],
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.alignment = Alignment(vertical="top", wrap_text=column in (1, 4, 8, 9))
        sheet.cell(row_index, 7).number_format = DATE_FORMAT

    last_row = max(header_row, header_row + len(rows))
    sheet.freeze_panes = f"A{header_row + 1}"
    widths = [28, 18, 16, 28, 18, 12, 16, 24, 27]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[header_row].height = 24
    sheet.auto_filter.ref = f"A{header_row}:I{last_row}"
    return workbook


def tape_measure_report_excel(request):
    if request.method != "GET":
        return JsonResponse({"error": "Metodă nepermisă."}, status=405)
    if not _admin_required(request):
        return JsonResponse({"error": "Nu ai dreptul să exporți raportul de rulete."}, status=403)
    start_date, end_date, error = _report_period(request)
    if error:
        return error
    rows = _filter_rows(
        tape_measure_report_rows(start_date, end_date),
        request.GET.get("q"),
    )
    workbook = _build_workbook(rows, start_date, end_date)
    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="raport_rulete_pana_la_{end_date:%Y-%m-%d}.xlsx"'
    response["X-Content-Type-Options"] = "nosniff"
    response["Cache-Control"] = "no-store"
    return response
